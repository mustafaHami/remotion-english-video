"""Read the user-facing Excel calendar. Writing stays explicit in build-workbook.mjs."""
from pathlib import Path
import argparse,json,datetime,sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
p=argparse.ArgumentParser();p.add_argument('--write',action='store_true');args=p.parse_args();root=Path(__file__).resolve().parents[2]
wb=load_workbook(root/'planning/calendrier-publications.xlsx',read_only=True,data_only=True);sheet=wb['Calendrier'];result=[]
assert [sheet.cell(5,i+1).value for i in range(15)]==['Date']+[f'Option {i}' for i in range(1,15)]
for values in sheet.iter_rows(min_row=6,max_col=15,values_only=True):
 if not values[0]:continue
 date=values[0].date().isoformat() if isinstance(values[0],datetime.datetime) else values[0].isoformat() if isinstance(values[0],datetime.date) else str(values[0])
 result.append(dict(date=date,options={str(i):str(v) for i,v in enumerate(values[1:],1) if v is not None and str(v).strip()}))
wb.close()
if args.write:(root/'planning/calendar.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
