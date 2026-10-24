"""Choix Codex/Claude, quotas, relance et nettoyage simulés : aucun agent réel n'est appelé."""
import copy
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from automation import cleanup, engines, planning, runner, worker
from automation.bridge import Application
from automation.store import Store
from model_helpers import QUOTAS, select_model

SESSION = '3f7c2a8e-1b2d-4c5e-9f00-112233445566'


def quota(engine, exhausted=False, status='ok', remaining=50):
    windows = [] if status != 'ok' else [{'label': '5 h', 'remaining': 0 if exhausted else remaining}]
    return {'engine': engine, 'status': status, 'exhausted': exhausted if status == 'ok' else None, 'windows': windows}


class EngineParsingTests(unittest.TestCase):
    def test_codex_limits_from_app_server(self):
        response = {'rateLimits': {'primary': {'usedPercent': 100, 'windowDurationMins': 10080}, 'secondary': None,
                                   'rateLimitReachedType': 'rate_limit_reached'}, 'ordinaryUsageAllowed': False}
        parsed = engines.parse_codex_limits(response)
        self.assertTrue(parsed['exhausted'])
        self.assertEqual(parsed['windows'], [{'label': 'semaine', 'remaining': 0}])
        ok = engines.parse_codex_limits({'rateLimits': {'primary': {'usedPercent': 32, 'windowDurationMins': 300}}})
        self.assertEqual((ok['exhausted'], ok['windows'][0]), (False, {'label': '5 h', 'remaining': 68}))
        self.assertIsNone(engines.parse_codex_limits({'rateLimits': {}})['exhausted'])

    def test_claude_limits_from_rate_limit_event(self):
        event = {'type': 'rate_limit_event', 'rate_limit_info': {'status': 'allowed', 'unifiedWindows': {
            'five_hour': {'utilization': 0.32, 'status': 'allowed'}, 'seven_day': {'utilization': 0.05, 'status': 'allowed'}}}}
        parsed = engines.parse_claude_limits(event)
        self.assertEqual(parsed['windows'], [{'label': '5 h', 'remaining': 68}, {'label': 'semaine', 'remaining': 95}])
        self.assertFalse(parsed['exhausted'])
        event['rate_limit_info']['status'] = 'rejected'
        self.assertTrue(engines.parse_claude_limits(event)['exhausted'])

    def test_quota_detected_in_agent_events(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'events.jsonl'
            path.write_text(json.dumps({'type': 'turn.failed', 'error': {'message': 'You’ve hit your usage limit.'}}) + '\n')
            self.assertTrue(engines.quota_from_events('codex', path))
            path.write_text(json.dumps({'type': 'error', 'message': 'Fichier introuvable'}) + '\n')
            self.assertFalse(engines.quota_from_events('codex', path))
            path.write_text(json.dumps({'type': 'result', 'is_error': True, 'api_error_status': 429}) + '\n')
            self.assertTrue(engines.quota_from_events('claude', path))

    def test_claude_not_configured_without_token(self):
        with patch.dict(os.environ, {'CREDENTIALS_DIRECTORY': '/nonexistent'}):
            self.assertEqual(engines.read_claude_quota()['status'], 'not_configured')


class RunnerTests(unittest.TestCase):
    def test_claude_runs_in_jail_with_exact_session_and_contract(self):
        job = {'id': 'a' * 12, 'revision': 1, 'option': 4, 'session': None}
        args = runner.claude_arguments(job, SESSION)
        self.assertEqual(args[0], 'bwrap')
        self.assertIn('--die-with-parent', args)
        command = args[args.index('--') + 1:]
        self.assertIn('--json-schema', command)
        self.assertEqual(command[command.index('--session-id') + 1], SESSION)
        self.assertEqual(command[command.index('--permission-mode') + 1], 'dontAsk')
        self.assertEqual(command[command.index('--model') + 1], engines.CLAUDE_MODEL)
        self.assertNotIn('bypassPermissions', command)
        settings = json.loads(command[command.index('--settings') + 1])
        self.assertIn('WebFetch', settings['permissions']['deny'])
        hidden = [args[i + 1] for i, a in enumerate(args[:args.index('--')]) if a == '--tmpfs']
        self.assertTrue(all(h in ['/tmp', *runner.HIDDEN] for h in hidden))
        resumed = runner.claude_arguments(job | {'session': SESSION}, SESSION, probe=True)
        self.assertEqual(resumed[0], engines.claude_binary())
        self.assertEqual(resumed[resumed.index('--resume') + 1], SESSION)

    def test_claude_events_give_session_usage_and_structured_result(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            events = directory / 'events.jsonl'
            events.write_text('\n'.join(json.dumps(e) for e in [
                {'type': 'system', 'subtype': 'init', 'session_id': SESSION},
                {'type': 'result', 'is_error': False, 'structured_output': {'ok': True},
                 'usage': {'input_tokens': 2, 'cache_creation_input_tokens': 10, 'cache_read_input_tokens': 5, 'output_tokens': 7}}]))
            session, usage = runner.session_usage(events, 'claude')
            self.assertEqual(session, SESSION)
            self.assertEqual(usage, {'input_tokens': 17, 'cached_input_tokens': 5, 'output_tokens': 7,
                                     'reasoning_output_tokens': 0, 'total_tokens': 24})
            runner.claude_result(events, directory)
            self.assertEqual(json.loads((directory / 'result.json').read_text()), {'ok': True})

    def test_handoff_prompt_resumes_from_files(self):
        self.assertIn('{previous}', runner.HANDOFF)
        self.assertIn('.claude/skills/create-english-video', runner.mission('claude'))
        self.assertIn('view_image', runner.mission('codex'))


class CleanupTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        for directory in ['content/episodes', 'output/options/option-4', 'public/audio/production', 'src/options/option-4']:
            (self.root / directory).mkdir(parents=True)
        (self.root / 'public/audio/production/provenance.json').write_text(json.dumps({'aaaaaaaaaaaaaaaa': {}}, indent=2))
        (self.root / 'public/audio/production/aaaaaaaaaaaaaaaa.wav').write_bytes(b'old')
        (self.root / 'content/episodes/old.json').write_text('{}')

    def tearDown(self):
        self.tmp.cleanup()

    def produce(self):
        before = cleanup.baseline(self.root)
        (self.root / 'content/episodes/new.json').write_text(json.dumps({'narration': [{'src': 'audio/production/bbbbbbbbbbbbbbbb.wav'}]}))
        folder = self.root / 'output/options/option-4/2026-10-24_20-00-00'
        folder.mkdir()
        (folder / 'option-4.mp4').write_bytes(b'mp4')
        for name in ['bbbbbbbbbbbbbbbb', 'cccccccccccccccc']:
            (self.root / f'public/audio/production/{name}.wav').write_bytes(b'new')
        meta = self.root / 'public/audio/production/provenance.json'
        data = json.loads(meta.read_text()) | {'bbbbbbbbbbbbbbbb': {}, 'cccccccccccccccc': {}}
        meta.write_text(json.dumps(data, indent=2))
        return cleanup.footprint(self.root, before)

    def test_footprint_lists_only_new_entries(self):
        self.assertEqual(self.produce(), [
            'content/episodes/new.json', 'output/options/option-4/2026-10-24_20-00-00',
            'provenance:bbbbbbbbbbbbbbbb', 'provenance:cccccccccccccccc',
            'public/audio/production/bbbbbbbbbbbbbbbb.wav', 'public/audio/production/cccccccccccccccc.wav'])

    def test_remove_files_keeps_shared_voice_and_old_files(self):
        entries = self.produce()
        # Une autre vidéo réutilise la voix ccc… : elle reste en place avec sa provenance.
        (self.root / 'content/episodes/other.json').write_text(json.dumps({'narration': [{'src': 'audio/production/cccccccccccccccc.wav'}]}))
        report = cleanup.remove_files(self.root, entries)
        self.assertFalse((self.root / 'content/episodes/new.json').exists())
        self.assertFalse((self.root / 'output/options/option-4/2026-10-24_20-00-00').exists())
        self.assertTrue((self.root / 'output/options/option-4').exists())
        self.assertFalse((self.root / 'public/audio/production/bbbbbbbbbbbbbbbb.wav').exists())
        self.assertTrue((self.root / 'public/audio/production/cccccccccccccccc.wav').exists())
        self.assertTrue((self.root / 'content/episodes/old.json').exists())
        provenance = json.loads((self.root / 'public/audio/production/provenance.json').read_text())
        self.assertEqual(set(provenance), {'aaaaaaaaaaaaaaaa', 'cccccccccccccccc'})
        self.assertEqual([s['path'] for s in report['skipped']], ['public/audio/production/cccccccccccccccc.wav'])

    def test_foreign_owner_links_and_invalid_entries_are_never_deleted(self):
        entries = self.produce()
        report = cleanup.remove_files(self.root, entries, uid=os.geteuid() + 1)
        self.assertEqual(report['deleted'], [])
        self.assertTrue((self.root / 'content/episodes/new.json').exists())
        target = self.root / 'outside.txt'
        target.write_text('keep')
        (self.root / 'content/episodes/link.json').symlink_to(target)
        cleanup.remove_files(self.root, ['content/episodes/link.json', '../outside.txt', 'planning/publications.json'])
        self.assertTrue(target.exists())
        self.assertFalse(cleanup.valid_entry('content/episodes/../../x.json'))

    def test_private_files_and_both_agent_sessions_removed(self):
        state, media = self.root / 'state', self.root / 'media'
        codex, claude = self.root / 'codex', self.root / 'claude'
        run_dir = state / 'runs' / ('a' * 12) / 'r1-x'
        run_dir.mkdir(parents=True)
        (run_dir / 'events.jsonl').write_text(json.dumps({'type': 'thread.started', 'thread_id': SESSION}))
        (media / ('a' * 12) / 'r1').mkdir(parents=True)
        (codex / 'sessions/2026/10/24').mkdir(parents=True)
        (codex / f'sessions/2026/10/24/rollout-{SESSION}.jsonl').write_text('{}')
        other = '11111111-2222-3333-4444-555555555555'
        (claude / 'projects/-srv-englizy-repository').mkdir(parents=True)
        (claude / f'projects/-srv-englizy-repository/{other}.jsonl').write_text('{}')
        (claude / 'projects/-srv-englizy-repository/keep.jsonl').write_text('{}')
        report = cleanup.remove_private('a' * 12, [other], state, media, codex, claude)
        self.assertFalse((state / 'runs' / ('a' * 12)).exists())
        self.assertFalse((media / ('a' * 12)).exists())
        self.assertFalse(list((codex / 'sessions').rglob('*.jsonl')))
        self.assertEqual([p.name for p in (claude / 'projects/-srv-englizy-repository').iterdir()], ['keep.jsonl'])
        self.assertIn('session:' + SESSION, report['deleted'])
        with self.assertRaises(ValueError):
            cleanup.remove_private('../etc', [], state, media, codex, claude)


class ModelFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.config = {'database': str(self.root / 'jobs.sqlite'), 'owner': 101, 'chat': 101,
                       'videoEnabled': True, 'mediaRoot': str(self.root)}
        self.app = Application(self.config)
        self.event = 0

    def tearDown(self):
        self.tmp.cleanup()

    def say(self, text, reply_to=None):
        self.event += 1
        message = {'text': text, 'from': {'id': 101}, 'chat': {'id': 101, 'type': 'private'}}
        if reply_to:
            message['reply_to_message'] = {'message_id': reply_to}
        return self.app.route('n8n', '/telegram', {'update_id': self.event, 'message': message})

    def click(self, data):
        self.event += 1
        return self.app.route('n8n', '/telegram', {'update_id': self.event, 'callback_query': {'id': str(self.event),
            'data': data, 'from': {'id': 101}, 'message': {'chat': {'id': 101, 'type': 'private'}}}})

    def messages(self):
        store = Store(self.config['database'])
        try:
            return [e['payload'] for e in store.pending()]
        finally:
            store.close()

    def last_text(self):
        return self.messages()[-1]['text']

    def store(self):
        return Store(self.config['database'])

    def open_menu(self, quotas=None):
        self.say('/modele')
        request = self.app.route('worker', '/worker/quota-request', {})['requestId']
        self.app.route('worker', '/worker/quotas', {'requestId': request, 'quotas': quotas or QUOTAS})
        return request

    def start(self, engine='codex'):
        self.open_menu()
        self.say('1' if engine == 'codex' else '2')
        self.say('/video')
        menu = self.messages()[-1]
        self.click(menu['keyboard']['rows'][0]['row']['buttons'][0]['additionalFields']['callback_data'])
        return self.app.route('worker', '/worker/claim', {'claimId': os.urandom(16).hex()})['job']

    def test_video_requires_model_every_time(self):
        self.say('/video')
        self.assertIn('oublié de choisir le modèle', self.last_text())
        job = self.start('claude')
        self.assertEqual(job['engine'], 'claude')
        store = self.store()
        self.assertEqual(store.db.execute("SELECT status FROM model_menus").fetchone()[0], 'consumed')
        store.db.execute("UPDATE jobs SET status='shared' WHERE id=?", (job['id'],))
        store.close()
        self.say('/video')
        self.assertIn('oublié de choisir le modèle', self.last_text())

    def test_menu_shows_quotas_and_refuses_exhausted_model(self):
        self.open_menu({'codex': quota('codex', exhausted=True), 'claude': quota('claude', status='not_configured')})
        menu = self.messages()[-1]
        self.assertIn('1. Codex', menu['text'])
        self.assertIn('épuisé', menu['text'])
        self.assertEqual([b['additionalFields']['callback_data'].split(':')[2] for b in menu['keyboard']['rows'][0]['row']['buttons']],
                         ['codex', 'claude'])
        self.say('1')
        self.assertIn('plus d’usage disponible', self.last_text())
        self.say('2')
        self.assertIn('pas encore configuré', self.last_text())
        self.say('/video')
        self.assertIn('oublié', self.last_text())

    def test_quota_read_timeout_publishes_unknown_and_allows_choice(self):
        self.say('/modele')
        self.say('1')
        self.assertIn('Lecture des quotas en cours', self.last_text())
        store = self.store()
        store.db.execute("UPDATE model_menus SET created='2020-01-01T00:00:00+00:00'")
        store.close()
        notes = self.app.route('n8n', '/outbox/claim', {'dispatchId': 'poll-1'})['notifications']
        self.assertIn('quota inconnu', notes[-1]['text'])
        self.assertEqual(notes[-1]['kind'], 'menu')
        self.say('2')
        self.assertIn('Modèle sélectionné : Claude', self.last_text())

    def test_quota_before_start_cancels_without_question(self):
        job = self.start()
        self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': False,
            'quotas': {'codex': quota('codex', exhausted=True), 'claude': quota('claude')}, 'footprint': []})
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'cancelling')
        store.close()
        self.assertIn('Aucune vidéo n’a été commencée', self.last_text())

    def test_fallback_yes_switches_agent_and_resumes_from_files(self):
        job = self.start()
        footprint = ['content/episodes/2026-10-24_20-00-00-option-1.json']
        self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': True,
            'quotas': {'codex': quota('codex', exhausted=True), 'claude': quota('claude', remaining=67)},
            'footprint': footprint, 'usage': {'input_tokens': 10, 'cached_input_tokens': 0, 'output_tokens': 5,
                                              'reasoning_output_tokens': 0, 'total_tokens': 15}, 'sessionId': SESSION})
        question = self.messages()[-1]
        self.assertEqual(question['kind'], 'menu')
        self.assertIn('Relancer avec Claude', question['text'])
        self.assertIn('67 % restant', question['text'])
        self.say('/modele')
        self.assertIn('Une vidéo reste à traiter', self.last_text())
        self.say('oui')
        claim = self.app.route('worker', '/worker/claim', {'claimId': 'b' * 32})['job']
        self.assertEqual((claim['id'], claim['engine'], claim['session']), (job['id'], 'claude', None))
        self.assertEqual(claim['handoff'], {'from': 'codex', 'reason': 'quota', 'files': footprint})
        store = self.store()
        self.assertEqual(set(store.usage_by_engine(job['id'])), {'codex'})
        store.close()

    def test_fallback_no_cancels_and_cleans_everything(self):
        job = self.start()
        self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': True,
            'quotas': {'codex': quota('codex', exhausted=True), 'claude': quota('claude')},
            'footprint': ['content/episodes/x.json'], 'sessionId': SESSION,
            'usage': {'input_tokens': 1, 'cached_input_tokens': 0, 'output_tokens': 1, 'reasoning_output_tokens': 0, 'total_tokens': 2}})
        fallback = next(m for m in self.messages() if m.get('keyboard', {}).get('rows') and 'Relancer' in m['text'])
        self.click(fallback['keyboard']['rows'][0]['row']['buttons'][1]['additionalFields']['callback_data'])
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'cancelling')
        self.assertFalse(store.db.execute("SELECT 1 FROM outbox WHERE id LIKE 'fallback:%' AND status='pending'").fetchone())
        store.close()
        task = self.app.route('worker', '/worker/claim', {'claimId': 'c' * 32})
        self.assertIsNone(task['job'])
        self.assertEqual(task['cleanup']['files'], ['content/episodes/x.json'])
        self.assertIn(SESSION, task['cleanup']['sessions'])
        again = self.app.route('worker', '/worker/claim', {'claimId': 'c' * 32})
        self.assertEqual(again['cleanup']['lease'], task['cleanup']['lease'])
        done = {'jobId': job['id'], 'lease': task['cleanup']['lease'], 'error': None,
                'report': {'deleted': ['content/episodes/x.json', 'session:' + SESSION], 'skipped': []}}
        self.app.route('worker', '/worker/cleanup-done', done)
        self.app.route('worker', '/worker/cleanup-done', done)
        store = self.store()
        cancelled = store.get(job['id'])
        self.assertEqual((cancelled['status'], cancelled['cancel_reason']), ('cancelled', 'fallback-declined'))
        self.assertIsNone(cancelled['session'])
        self.assertIsNotNone(cancelled['purged_at'])
        for table in ['questions', 'revisions', 'job_files', 'codex_usage', 'preview_media']:
            self.assertEqual(store.db.execute(f'SELECT COUNT(*) FROM {table} WHERE job=?', (job['id'],)).fetchone()[0], 0, table)
        self.assertEqual([r[0] for r in store.db.execute('SELECT action FROM audit WHERE job=?', (job['id'],))], ['cancelled-cleaned'])
        self.assertFalse(store.outstanding(101, 101))
        store.close()
        self.assertIn('annulée', self.last_text())
        self.say('/modele')
        self.assertIn('Lecture des quotas', self.last_text())

    def test_both_agents_exhausted_cancel_directly(self):
        job = self.start()
        self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': True,
            'quotas': {'codex': quota('codex', exhausted=True), 'claude': quota('claude', exhausted=True)}, 'footprint': []})
        store = self.store()
        self.assertEqual(store.get(job['id'])['cancel_reason'], 'quota-all')
        store.close()
        self.assertIn('Codex et Claude n’ont plus d’usage', self.last_text())
        repeated = self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': True,
            'quotas': {'codex': quota('codex', exhausted=True), 'claude': quota('claude', exhausted=True)}, 'footprint': []})
        self.assertTrue(repeated['alreadyRecorded'])

    def test_cancel_rules_and_failed_cleanup_retry(self):
        job = self.start()
        self.say(f"/annuler {job['id']}")
        self.assertIn('Annulation impossible maintenant', self.last_text())
        self.app.route('worker', '/worker/failure', {'jobId': job['id'], 'lease': job['lease'], 'footprint': ['content/episodes/y.json']})
        failure = self.messages()[-1]
        self.assertEqual(failure['keyboard']['rows'][0]['row']['buttons'][0]['additionalFields']['callback_data'], f"cancel:{job['id']}")
        self.say(f"/annuler {job['id']}")
        task = self.app.route('worker', '/worker/claim', {'claimId': 'd' * 32})['cleanup']
        self.assertEqual(task['files'], ['content/episodes/y.json'])
        self.app.route('worker', '/worker/cleanup-done', {'jobId': job['id'], 'lease': task['lease'], 'report': None, 'error': 'OSError: busy'})
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'cleanup_failed')
        store.close()
        self.say(f"/annuler {job['id']}")
        task = self.app.route('worker', '/worker/claim', {'claimId': 'e' * 32})['cleanup']
        self.app.route('worker', '/worker/cleanup-done', {'jobId': job['id'], 'lease': task['lease'],
                                                          'report': {'deleted': [], 'skipped': []}, 'error': None})
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'cancelled')
        store.close()

    def test_invalid_worker_payloads_are_refused(self):
        job = self.start()
        for footprint in [['../../etc/passwd'], ['planning/publications.json'], 'content/episodes/x.json']:
            with self.assertRaises(ValueError):
                self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': job['lease'], 'started': True,
                    'quotas': QUOTAS, 'footprint': footprint})
        with self.assertRaises(ValueError):
            self.app.route('worker', '/worker/quotas', {'requestId': 'x', 'quotas': {'codex': {'engine': 'codex', 'windows': 'x'}}})
        with self.assertRaises(ValueError):
            self.app.route('worker', '/worker/quota-exhausted', {'jobId': job['id'], 'lease': 'wrong', 'started': True,
                'quotas': QUOTAS, 'footprint': []})

    def test_worker_cycle_reports_quota_and_runs_cleanup_after_restart(self):
        job = self.start()
        store = self.store()
        store.db.execute("UPDATE jobs SET status='queued',lease=NULL WHERE id=?", (job['id'],))
        store.close()
        config = {'stateDir': str(self.root / 'worker'), 'bridge': 'http://invalid.test', 'videoEnabled': True,
                  'mediaRoot': str(self.root / 'media')}
        transport = lambda url, token, path, body: self.app.route('worker', path, body)
        reads = iter([quota('codex'), quota('codex', exhausted=True), quota('claude')])
        with patch.object(worker, 'request', side_effect=transport), \
                patch.object(worker, 'run', side_effect=runner.QuotaExhausted('épuisé')), \
                patch.object(worker.engines, 'read_quota', side_effect=lambda e: next(reads)), \
                patch.object(worker.cleanup, 'baseline', return_value=[]), \
                patch.object(worker.cleanup, 'footprint', return_value=['content/episodes/z.json']):
            worker.cycle(config, 'fake')
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'needs_fallback')
        self.assertEqual(store.files(job['id']), ['content/episodes/z.json'])
        store.close()
        self.say('non')
        calls = []
        with patch.object(worker, 'request', side_effect=transport), \
                patch.object(worker, 'clean', side_effect=lambda c, t: calls.append(t['jobId']) or {'deleted': [], 'skipped': []}):
            # Panne simulée après la réservation : l'état durable relance le même nettoyage.
            state_file = Path(config['stateDir']) / 'active.json'
            response = self.app.route('worker', '/worker/claim', {'claimId': 'f' * 32})
            worker.save(state_file, {'phase': 'cleaning', 'claimId': 'f' * 32, 'cleanup': response['cleanup']})
            worker.cycle(config, 'fake')
        self.assertEqual(calls, [job['id']])
        store = self.store()
        self.assertEqual(store.get(job['id'])['status'], 'cancelled')
        store.close()


class PlanningRemovalTests(unittest.TestCase):
    def test_cancelled_job_leaves_excel_and_production_index(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'planning').mkdir()
            path = root / 'planning/calendrier-publications.xlsx'
            shutil.copyfile(Path(__file__).resolve().parents[1] / 'planning/calendrier-publications.xlsx', path)
            (root / 'planning/publications.json').write_text('[]')
            database = root / 'jobs.sqlite'
            store = Store(database)
            job = store.create(101, 101, 1)['jobId']
            store.close()
            record = {'id': job, 'date': '2026-10-24', 'option': 1, 'status': 'generating', 'revision': 1,
                      'content': [], 'summary': '', 'platforms': [], 'publicationUrls': [],
                      'source': 'Suivi de fabrication ; aucune publication confirmée.', 'episode': None, 'validationSource': None}
            with patch.object(planning, 'read_records', side_effect=lambda *a: copy.deepcopy([record])):
                planning.sync(database, root, root / 'backup')
            wb = load_workbook(path)
            self.assertTrue(any(r[4].value == job for r in wb['Journal'].iter_rows(min_row=6)))
            self.assertIn('[Suivi ' + job + ']', wb['Calendrier'].cell(row=next(
                i for i in range(6, wb['Calendrier'].max_row + 1) if '[Suivi' in str(wb['Calendrier'].cell(i, 2).value or '')), column=2).value)
            store = Store(database)
            store.db.execute("UPDATE jobs SET status='cancelled',cancel_reason='user-request',purged_at='2026-10-24T20:00:00+00:00' WHERE id=?", (job,))
            store.close()
            with patch.object(planning, 'read_records', return_value=[]):
                planning.sync(database, root, root / 'backup')
            wb = load_workbook(path)
            self.assertFalse(any(r[4].value == job for r in wb['Journal'].iter_rows(min_row=6)))
            self.assertFalse(any('[Suivi ' + job in str(c.value or '') for row in wb['Calendrier'].iter_rows() for c in row))
            index = json.loads((root / 'planning/production.json').read_text())
            self.assertNotIn(job, [v['id'] for v in index['videos']])


if __name__ == '__main__':
    unittest.main()
