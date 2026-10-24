import datetime
import copy
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from automation.store import Store
from automation.bridge import Application
from automation import planning, retention, tracking


class TrackingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.database = self.root/'jobs.sqlite'
        self.store = Store(self.database)
        self.config = {'database': str(self.database), 'owner':101, 'chat':101, 'videoEnabled':True, 'dailySummaryTime':'09:00'}
        self.app = Application(self.config)

    def tearDown(self):
        self.store.close(); self.tmp.cleanup()

    def ready(self):
        job = self.store.create(101,101,1)
        self.store.db.execute("UPDATE jobs SET status='pending_validation' WHERE id=?", (job['jobId'],))
        media = {'originalSha256':'a'*64, 'sha256':'a'*64}
        self.store.db.execute('INSERT INTO preview_media VALUES(?,?,?)', (job['jobId'],1,json.dumps(media)))
        self.store.emit('ready:'+job['jobId']+':1',job['jobId'],{'kind':'ready','revision':1})
        self.store.mark_delivery('ready:'+job['jobId']+':1','sent',{'messageId':10,'chatId':101})
        return job['jobId']

    def test_validation_is_exact_and_does_not_authorize_publication(self):
        job = self.ready()
        with self.assertRaises(PermissionError): self.store.validate_preview(job,1,102,101)
        with self.assertRaises(ValueError): self.store.validate_preview(job,2,101,101)
        self.store.validate_preview(job,1,101,101)
        self.store.validate_preview(job,1,101,101)
        self.assertEqual(self.store.get(job)['status'],'validated')
        self.assertIsNone(self.store.get(job)['approval'])
        self.assertIsNone(self.store.get(job)['closed_at'])
        with self.assertRaises(ValueError): self.store.confirm_shared(job,1,101,101,{'source':'telegram-receipt'})

    def test_outstanding_video_blocks_new_menu_and_stale_choice(self):
        menu = self.app.options(self.store)
        choice = menu['keyboard']['rows'][0]['row']['buttons'][0]['additionalFields']['callback_data'].split(':')[1]
        job = self.ready()
        for status in ['pending_validation','validated','needs_input','failed','generating']:
            self.store.db.execute('UPDATE jobs SET status=? WHERE id=?',(status,job))
            with self.assertRaises(ValueError): self.app.options(self.store)
            with self.assertRaises(ValueError): self.app.choose(self.store,choice,2)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM jobs').fetchone()[0],1)

    def test_external_workbook_record_also_blocks_creation(self):
        p=self.root/'production.json';p.write_text(json.dumps({'videos':[{'id':'manual-video','option':4,'status':'validated'}]}))
        app=Application(self.config|{'planningIndex':str(p)})
        with self.assertRaises(ValueError):app.options(self.store)

    def test_telegram_shared_requires_explicit_confirmation_and_correct_version(self):
        job=self.ready();self.store.validate_preview(job,1,101,101)
        def update(event,text):
            return {'update_id':event,'message':{'text':text,'from':{'id':101},'chat':{'id':101,'type':'private'}}}
        self.app.route('n8n','/telegram',update(1,f'/partage {job} 1 Instagram'))
        self.assertEqual(self.store.get(job)['status'],'validated')
        self.app.route('n8n','/telegram',update(2,f'/partage {job} 2 CONFIRME Instagram'))
        self.assertEqual(self.store.get(job)['status'],'validated')
        result=self.app.route('n8n','/telegram',update(3,f'/partage {job} 1 CONFIRME Instagram'))
        self.assertEqual(result['status'],'shared')
        j=self.store.get(job);self.assertEqual(j['closed_at'],j['shared_at'])
        self.assertEqual(json.loads(j['sharing'])['platforms'],['Instagram'])
        self.assertEqual(self.app.route('n8n','/telegram',update(3,f'/partage {job} 1 CONFIRME Instagram')),result)
        self.assertEqual(tracking.outstanding(self.store,self.config),[])

    def test_daily_summary_timezone_and_restart_dedup(self):
        self.ready()
        # Paris : UTC+2 en été, UTC+1 en hiver ; jamais envoyer avant 09h.
        for date,hour in [('2026-07-10',7),('2026-12-10',8)]:
            due=datetime.datetime.fromisoformat(date).replace(hour=hour,tzinfo=datetime.timezone.utc)
            self.assertTrue(tracking.daily_summary(self.store,self.config,due-datetime.timedelta(seconds=1))['notDue'])
            tracking.daily_summary(self.store,self.config,due)
            other=Store(self.database)
            self.assertTrue(tracking.daily_summary(other,self.config,due)['alreadyRecorded']);other.close()
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM outbox WHERE id LIKE 'daily-summary:%'").fetchone()[0],2)

    def test_incomplete_partage_explains_declaration_without_changing_state(self):
        job=self.ready();self.store.validate_preview(job,1,101,101)
        for event,text in enumerate(['/partage',f'/partage {job}',f'/partage {job} 1',f'/partage {job} 1 Instagram'],start=100):
            update={'update_id':event,'message':{'text':text,'from':{'id':101},'chat':{'id':101,'type':'private'}}}
            self.app.route('n8n','/telegram',update)
            payload=json.loads(self.store.db.execute('SELECT payload FROM outbox WHERE id=?',('event:'+str(event),)).fetchone()[0])
            self.assertIn('ne lance pas la publication',payload['text'])
            self.assertIn('IDENTIFIANT VERSION CONFIRME',payload['text'])
            current=self.store.get(job)
            self.assertEqual(current['status'],'validated')
            self.assertIsNone(current['shared_at']);self.assertIsNone(current['closed_at']);self.assertIsNone(current['approval'])
            self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM jobs').fetchone()[0],1)

    def test_excel_preserves_existing_briefs_and_legacy_rows_and_is_idempotent(self):
        (self.root/'planning').mkdir()
        path=self.root/'planning/calendrier-publications.xlsx'
        shutil.copyfile(Path(__file__).resolve().parents[1]/'planning/calendrier-publications.xlsx',path)
        wb=load_workbook(path);before={c.coordinate:c.value for row in wb['Journal'].iter_rows(min_row=6,max_row=7) for c in row}
        record={'id':'aaaaaaaaaaaa','date':'2026-10-24','option':1,'status':'validated','revision':1,
                'content':['=FORMULA MUST STAY TEXT','I work. I worked.'], 'platforms':[],'publicationUrls':[],
                'source':'Validation de la prévisualisation, non publiée','episode':None,'validationSource':'User'}
        wb['Calendrier']['B10']='Brief humain à préserver';wb.save(path)
        with patch.object(planning,'read_records',side_effect=lambda *args: copy.deepcopy([record])):
            self.assertTrue(planning.sync(self.database,self.root,self.root/'backup')['workbookUpdated'])
            first=path.read_bytes()
            self.assertFalse(planning.sync(self.database,self.root,self.root/'backup')['workbookUpdated'])
            self.assertEqual(first,path.read_bytes())
        out=load_workbook(path)
        self.assertEqual(before,{c.coordinate:c.value for row in out['Journal'].iter_rows(min_row=6,max_row=7) for c in row})
        self.assertTrue(out['Calendrier']['B10'].value.startswith('Brief humain à préserver'))
        found=[r for r in out['Journal'].iter_rows(min_row=6) if r[4].value==record['id']]
        self.assertEqual(len(found),1);self.assertEqual(found[0][2].value,'validated');self.assertEqual(found[0][3].data_type,'s')
        found[0][7].value = 'Note humaine à préserver';out.save(path)
        edited = path.read_bytes()
        with patch.object(planning,'read_records',side_effect=lambda *args: copy.deepcopy([record])):
            with self.assertRaises(ValueError): planning.sync(self.database,self.root,self.root/'backup')
        self.assertEqual(edited,path.read_bytes())

    def test_retention_never_deletes_open_video_and_respects_10_15_day_boundaries(self):
        job=self.ready();self.store.validate_preview(job,1,101,101)
        moment=datetime.datetime(2026,11,20,tzinfo=datetime.timezone.utc)
        # Ancienne vidéo validée sans partage : aucune suppression.
        self.store.db.execute('UPDATE jobs SET updated=? WHERE id=?',('2020-01-01T00:00:00+00:00',job))
        self.assertEqual(retention.apply(self.store,self.root,self.root/'media',self.root/'backup',moment)['mediaDeleted'],[])
        self.store.confirm_shared(job,1,101,101,{'source':'explicit-user-confirmation','allDestinationsConfirmed':True,'confirmation':'Test simulé : publication terminée'})
        at=(moment-datetime.timedelta(days=10)).isoformat()
        self.store.db.execute('UPDATE jobs SET shared_at=?,closed_at=? WHERE id=?',(at,at,job))
        video=self.root/'output/options/option-1/2026-10-01/option-1.mp4';video.parent.mkdir(parents=True);video.write_bytes(b'mp4-test')
        spool=self.root/'media';preview=spool/job/'r1/preview.mp4';preview.parent.mkdir(parents=True);preview.write_bytes(video.read_bytes())
        h=hashlib.sha256(video.read_bytes()).hexdigest()
        self.store.db.execute('UPDATE preview_media SET media=? WHERE job=?',(json.dumps({'sha256':h,'originalSha256':h}),job))
        self.store.db.execute('UPDATE jobs SET result=? WHERE id=?',(json.dumps({'artifacts':{'video':str(video)}}),job))
        public=self.root/'public-media';public.mkdir()
        public_copy=public/('a'*64+'.mp4');public_copy.write_bytes(video.read_bytes())
        self.store.db.execute('INSERT INTO buffer_media VALUES(?,?,?,?,?)',(job,1,public_copy.name,h,'https://example.com/video.mp4'))
        self.store.db.execute("INSERT INTO buffer_publications(job,revision,channel,service,post_id,expected,edit_input,status,response) VALUES(?,1,'channel','instagram','post','{\"text\":\"caption\"}','{\"text\":\"caption\"}','sent',?)",
                              (job,json.dumps({'status':'sent','sentAt':at,'text':'full caption'})))
        before=moment-datetime.timedelta(seconds=1)
        self.assertFalse(retention.apply(self.store,self.root,spool,self.root/'backup',before,public)['mediaDeleted'])
        preview.write_bytes(b'modified')
        with self.assertRaises(ValueError):retention.apply(self.store,self.root,spool,self.root/'backup',moment,public)
        self.assertTrue(video.exists())
        preview.write_bytes(video.read_bytes())
        public_copy.write_bytes(b'modified-public-copy')
        with self.assertRaises(ValueError):retention.apply(self.store,self.root,spool,self.root/'backup',moment,public)
        self.assertTrue(video.exists())
        public_copy.write_bytes(video.read_bytes())
        self.assertEqual(retention.apply(self.store,self.root,spool,self.root/'backup',moment,public)['mediaDeleted'],[job])
        self.assertFalse(public_copy.exists())
        self.assertFalse(video.exists());self.assertFalse(preview.exists())
        self.assertIsNone(self.store.get(job)['purged_at'])
        result=retention.apply(self.store,self.root,spool,self.root/'backup',moment+datetime.timedelta(days=5))
        self.assertEqual(result['detailsPurged'],[job]);self.assertIsNone(self.store.get(job)['result'])
        self.assertEqual(self.store.get(job)['status'],'shared')
        self.assertTrue(self.store.get(job)['sharing'])
        publication=self.store.db.execute('SELECT expected,edit_input,response FROM buffer_publications WHERE job=?',(job,)).fetchone()
        self.assertEqual(publication['expected'],'{}');self.assertEqual(publication['edit_input'],'{}')
        self.assertEqual(json.loads(publication['response']),{'status':'sent','sentAt':at,'externalLink':None})


if __name__ == '__main__': unittest.main()
