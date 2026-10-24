"""Projection du suivi vers Excel, sans transformer un rendu en publication."""
from copy import copy
import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from .store import connect_as_owner

HEADERS = ['Date', 'Option', 'Statut', 'Contenu utilisé', 'Identifiant', 'Plateformes', 'Liens', 'Source']


def read_records(database, root):
    root = Path(root).resolve()
    db = connect_as_owner(database, readonly=True)
    db.row_factory = sqlite3.Row
    try:
        records = []
        for row in db.execute("SELECT * FROM jobs WHERE kind='video' AND purged_at IS NULL ORDER BY rowid"):
            j = dict(row)
            result = json.loads(j['result']) if j['result'] else {}
            artifacts = result.get('artifacts', {}) if result.get('revision') == j['revision'] else {}
            created = db.execute("SELECT at FROM audit WHERE job=? AND action='created' ORDER BY id LIMIT 1", (j['id'],)).fetchone()
            day = datetime.datetime.fromisoformat(created[0] if created else j['updated']).astimezone(ZoneInfo('Europe/Paris')).date().isoformat()
            content, episode_relative = [], None
            if artifacts.get('episode'):
                value = artifacts['episode'].removeprefix('/srv/englizy-repository/')
                episode = (root / value).resolve(strict=True)
                if not episode.is_relative_to(root / 'content/episodes'):
                    raise ValueError('Épisode hors périmètre du suivi.')
                data = json.loads(episode.read_text())
                if data['option'] != j['option']:
                    raise ValueError('Option du suivi incohérente.')
                content = [c['text'] for c in data['narration'] if c['role'] == 'lesson']
                episode_relative = str(episode.relative_to(root))
            sharing = json.loads(j['sharing']) if j['sharing'] else {}
            if j['status'] == 'shared' and (not sharing.get('confirmation') or sharing.get('allDestinationsConfirmed') is not True):
                raise ValueError('Publication non confirmée : historique inchangé.')
            records.append({'id': j['id'], 'date': day, 'option': j['option'], 'status': j['status'],
                'revision': j['revision'], 'content': content, 'summary': result.get('summary', ''),
                'episode': episode_relative, 'validatedAt': j['validated_at'], 'sharedAt': j['shared_at'],
                'closedAt': j['closed_at'], 'platforms': sharing.get('platforms', []),
                'publicationUrls': sharing.get('urls', []),
                'source': ('Publication terminée confirmée : ' + sharing.get('confirmation', '') if j['status'] == 'shared'
                           else 'Suivi de fabrication ; aucune publication confirmée.'),
                'validationSource': json.loads(j['review'])['source'] if j['review'] else None})
        return records
    finally:
        db.close()


def cancelled_ids(database):
    """Tâches annulées et nettoyées : elles disparaissent du suivi lisible."""
    db = connect_as_owner(database, readonly=True)
    try:
        return {r[0] for r in db.execute("SELECT id FROM jobs WHERE kind='video' AND status='cancelled' AND cancel_reason IS NOT NULL")}
    finally:
        db.close()


def remove_blocks(calendar, removed):
    """Retire le bloc [Suivi] d'une tâche annulée sans toucher aux briefs voisins."""
    for identifier in removed:
        start, end = '[Suivi ' + identifier + ']', '[/Suivi ' + identifier + ']'
        for row in calendar.iter_rows(min_row=6):
            for cell in row:
                if not isinstance(cell.value, str) or start not in cell.value:
                    continue
                if cell.data_type == 'f':
                    raise ValueError('Cellule de planning contenant une formule : réconciliation requise.')
                pattern = r'(\n\n)?' + re.escape(start) + r'.*?' + re.escape(end)
                if not re.search(pattern, cell.value, flags=re.S):
                    raise ValueError('Bloc de suivi incomplet dans le Calendrier.')
                set_value(cell, re.sub(pattern, '', cell.value, flags=re.S).strip())


def set_value(cell, value):
    if value == '':
        value = None
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError('Texte trop long pour une cellule Excel ; aucune troncature autorisée.')
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # Le contenu pédagogique ne doit jamais devenir une formule.


def journal_values(r):
    source = r['source']
    if r.get('validationSource'):
        source += '\nValidation : ' + r['validationSource']
    source += f"\nVersion {r['revision']}" + ('\nÉpisode : ' + r['episode'] if r.get('episode') else '')
    return [datetime.datetime.fromisoformat(r['date']), r['option'], r['status'],
            '\n'.join(r['content']) or r.get('summary', ''), r['id'],
            ', '.join(r['platforms']), '\n'.join(r['publicationUrls']), source]


def update_workbook(workbook, records, previous=None, removed=()):
    journal, calendar = workbook['Journal'], workbook['Calendrier']
    if [journal.cell(5, c).value for c in range(1, 9)] != HEADERS:
        raise ValueError('En-têtes du Journal modifiés : réconciliation requise.')
    if [calendar.cell(5, c).value for c in range(1, 16)] != ['Date'] + [f'Option {n}' for n in range(1, 15)]:
        raise ValueError('En-têtes du Calendrier modifiés : réconciliation requise.')
    ids = {}
    for row in range(6, journal.max_row + 1):
        identifier = journal.cell(row, 5).value
        if identifier:
            if identifier in ids:
                raise ValueError('Identifiant dupliqué dans le Journal.')
            ids[identifier] = row
    for r in previous or []:
        if r.get('external') or r['id'] not in ids:
            continue
        observed = [journal.cell(ids[r['id']], c).value for c in range(1, 9)]
        expected = [None if v == '' else v for v in journal_values(r)]
        if observed != expected:
            raise ValueError('Ligne de suivi modifiée dans Excel : réconciliation requise pour ' + r['id'])
    if removed:
        rows = sorted(((ids.pop(i), i) for i in list(removed) if i in ids), reverse=True)
        for row, _ in rows:
            journal.delete_rows(row)
        for identifier in list(ids):
            ids[identifier] -= sum(1 for row, _ in rows if row < ids[identifier])
        remove_blocks(calendar, removed)
    for r in records:
        row = ids.get(r['id'])
        if row is None:
            row = max([5, *ids.values()]) + 1
            ids[r['id']] = row
            for col in range(1, 9):
                journal.cell(row, col)._style = copy(journal.cell(6, col)._style)
            journal.row_dimensions[row].height = 200
        values = journal_values(r)
        for col, value in enumerate(values, 1):
            set_value(journal.cell(row, col), value)
            journal.cell(row, col).alignment = Alignment(wrap_text=True, vertical='top')
        journal.cell(row, 1).number_format = 'dd/mm/yyyy'
        journal.cell(row, 3).fill = PatternFill('solid', fgColor={'validated':'D4F5E3','shared':'A9D0F5'}.get(r['status'],'FFF2CC'))
        journal.row_dimensions[row].height = min(409, max(200, (len(r['content']) + 1) * 15))
        date_row = None
        for i in range(6, calendar.max_row + 1):
            value = calendar.cell(i, 1).value
            date = value.date().isoformat() if isinstance(value, datetime.datetime) else str(value)
            if date == r['date']:
                date_row = i; break
        if date_row is None:
            date_row = calendar.max_row + 1
            calendar.cell(date_row, 1, datetime.datetime.fromisoformat(r['date'])).number_format = 'dd/mm/yyyy'
        cell = calendar.cell(date_row, r['option'] + 1)
        if cell.data_type == 'f':
            raise ValueError('Cellule de planning contenant une formule : réconciliation requise.')
        existing = str(cell.value or '')
        start, end = '[Suivi ' + r['id'] + ']', '[/Suivi ' + r['id'] + ']'
        block = start + '\n' + r['status'] + '\n' + ('\n'.join(r['content']) or r.get('summary', '')) + '\n' + end
        if start in existing:
            pattern = re.escape(start) + r'.*?' + re.escape(end)
            if not re.search(pattern, existing, flags=re.S):
                raise ValueError('Bloc de suivi incomplet dans le Calendrier.')
            value = re.sub(pattern, lambda _: block, existing, flags=re.S)
        else:
            value = existing + ('\n\n' if existing else '') + block
        set_value(cell, value)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        calendar.row_dimensions[date_row].height = max(calendar.row_dimensions[date_row].height or 48,
                                                       min(409, (len(r['content']) + 4) * 15))
    journal.column_dimensions['C'].width = 24


def sync(database, root, backup_dir):
    root, backup_dir = Path(root), Path(backup_dir)
    workbook_path = root / 'planning/calendrier-publications.xlsx'
    index_path = root / 'planning/production.json'
    before = workbook_path.read_bytes()
    wb = load_workbook(workbook_path)
    records = read_records(database, root)
    # Conserver les fiches minimales après purge des détails en base, sauf les tâches annulées.
    old = json.loads(index_path.read_text())['videos'] if index_path.exists() else []
    cancelled = cancelled_ids(database)
    current_ids = {r['id'] for r in records}
    records += [r for r in old if r['id'] not in current_ids and not r.get('external') and r['id'] not in cancelled]
    update_workbook(wb, records, previous=old, removed={r['id'] for r in old if r['id'] in cancelled and not r.get('external')})
    # Le Journal peut aussi contenir des productions humaines non gérées par le bot.
    indexed = {r['id'] for r in records}
    for row in wb['Journal'].iter_rows(min_row=6, max_col=8, values_only=True):
        day, option, status, content, identifier, platforms, urls, source = row
        if not identifier or identifier in indexed:
            continue
        normalized = {'Publié':'shared', 'published':'shared', 'pending validation':'pending_validation'}.get(status, status)
        records.append({'id': str(identifier), 'date': day.date().isoformat() if isinstance(day, datetime.datetime) else str(day),
                        'option': option, 'status': normalized, 'revision': 1, 'content': str(content or '').splitlines(),
                        'platforms': str(platforms or '').splitlines(), 'publicationUrls': str(urls or '').splitlines(),
                        'source': source or '', 'external': True})
    # Les anciennes publications restent intactes ; leur statut publié équivaut à shared.
    data = {'schemaVersion': 1, 'videos': sorted(records, key=lambda r: (r['date'], r['id']))}
    serialized = json.dumps(data, ensure_ascii=False, indent=2) + '\n'
    # Comparer les valeurs et styles avant d'écrire : pas de réécriture toutes les minutes.
    original = load_workbook(workbook_path)
    changed = any([(c.coordinate, c.value, str(c._style)) for row in s for c in row] !=
                  [(c.coordinate, c.value, str(c._style)) for row in original[s.title] for c in row] for s in wb)
    changed = changed or any(
        {k:v.height for k,v in s.row_dimensions.items()} != {k:v.height for k,v in original[s.title].row_dimensions.items()}
        or {k:v.width for k,v in s.column_dimensions.items()} != {k:v.width for k,v in original[s.title].column_dimensions.items()}
        for s in wb)
    backup_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
    if changed:
        digest = hashlib.sha256(before).hexdigest()
        backup = backup_dir / (digest + '.xlsx')
        if not backup.exists():
            backup.write_bytes(before)
        with tempfile.NamedTemporaryFile(dir=workbook_path.parent, suffix='.xlsx', delete=False) as f:
            temporary = Path(f.name)
        try:
            wb.save(temporary)
            check = load_workbook(temporary)
            assert check.sheetnames == original.sheetnames
            if workbook_path.read_bytes() != before:
                raise ValueError('Classeur modifié pendant la synchronisation ; remplacement refusé.')
            temporary.chmod(workbook_path.stat().st_mode & 0o777)
            temporary.replace(workbook_path)
        finally:
            temporary.unlink(missing_ok=True)
    if not index_path.exists() or index_path.read_text() != serialized:
        temporary = index_path.with_suffix('.json.tmp')
        temporary.write_text(serialized); temporary.chmod(0o644); temporary.replace(index_path)
    # Compatibilité du registre historique : published y signifie shared/clos.
    history_path = root / 'planning/publications.json'
    shared = [r for r in records if r['status'] == 'shared' and not r.get('external') and r.get('sharedAt')]
    if shared:
        before_history = history_path.read_bytes()
        history = json.loads(before_history)
        for r in shared:
            entry = {'id':r['id'], 'date':datetime.datetime.fromisoformat(r['sharedAt']).astimezone(ZoneInfo('Europe/Paris')).date().isoformat(),
                     'option':r['option'], 'status':'published', 'content':r['content'], 'platforms':r['platforms'],
                     'publicationUrls':r['publicationUrls'], 'source':r['source']}
            old_entry = next((item for item in history if item['id']==r['id']), None)
            if old_entry and old_entry != entry:
                raise ValueError('Publication existante différente ; réconciliation requise.')
            if not old_entry:
                history.append(entry)
        content = json.dumps(history,ensure_ascii=False,indent=2)+'\n'
        if content.encode() != before_history:
            if history_path.read_bytes() != before_history:
                raise ValueError('Historique modifié pendant la synchronisation.')
            target=history_path.with_suffix('.json.tmp');target.write_text(content);target.chmod(0o644);target.replace(history_path)
    return {'workbookUpdated': changed, 'videos': len(records)}
